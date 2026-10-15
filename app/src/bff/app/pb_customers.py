"""Customer 360, financial wellness, next-best action and outreach - deterministic core.

Every number comes from the governed Fabric IQ tables (bundled in pb_data). The
rules that gate an action - consent (RP-011), vulnerability (RP-012),
affordability (RP-013) - are applied here in code, before any AI is involved.
The compliance reviewer agent only adds judgement (tone, implied promises) on
top of these deterministic findings; it can block, never approve on its own.
SYNTHETIC demo data.
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import date, timedelta
from typing import Any

from .pb_data import PB

AS_OF = date.fromisoformat(PB["as_of"])
CUST = PB["customers"]
REN_BY_CUST = {r["customer_id"]: r for r in PB["renewals"].values()}
RISK = PB["risk"]
CAP = {r["customer_id"]: r for r in PB["credit_capacity"]}
WELL = {r["customer_id"]: r for r in PB["financial_wellness"]}
HH = {r["primary_customer_id"]: r for r in PB["households"]}
_by = lambda rows: (lambda d: [d[k].append(r) for r in rows for k in [r["customer_id"]]] and d)(defaultdict(list))  # noqa: E731
PRODUCTS = _by(PB["customer_products"])
SIGNALS = _by(PB["financial_signals_monthly"])
EVENTS = _by(PB["life_events"])
CONSENT = _by(PB["consent_preferences"])
NBA = _by(PB["next_best_actions"])
TDS_LIMIT, GDS_LIMIT = PB["valuation"]["tds_limit_pct"], PB["valuation"]["gds_limit_pct"]

BANNED = [
    (r"\bguarantee(d)?\b", "RP-016", "Prohibited claim: 'guaranteed'"),
    (r"\bpre-?approved\b", "RP-013", "Implies approval before adjudication: 'pre-approved'"),
    (r"\brisk[- ]free\b", "RP-016", "Prohibited claim: 'risk-free'"),
    (r"\bbest rate\b", "RP-016", "Unsupported superlative: 'best rate'"),
    (r"\b(act now|limited time|don't miss|last chance|expires today)\b", "RP-016", "Urgency / scarcity framing"),
    (r"\b(attrition|wellness (score|band)|vulnerab\w*|risk score|TDS|GDS|approval tier)\b", "RP-016",
     "Internal-only data exposed to the client"),
    (r"\b(RBC|TD|CIBC|BMO|Scotiabank|National Bank|Desjardins|Tangerine|Simplii|EQ Bank)\b", "RP-016",
     "Competitor or bank named in client-facing copy"),
]


def _consent(cid: str) -> dict[str, str]:
    return {c["channel"]: c["consent_status"] for c in CONSENT[cid]}


def _flags(cid: str) -> list[str]:
    w = WELL[cid]
    f = []
    if w["vulnerability_flag"] == "Y":
        f.append("Support first")
    if float(w["competitor_outflow_share_pct"]) >= 10:
        f.append("Deposit flight")
    if float(CAP[cid]["tds_renewal_pct"]) > TDS_LIMIT:
        f.append("Above TDS guideline")
    if any(e["implication"] == "Deepen" for e in EVENTS[cid]):
        f.append("Life event")
    return f


def _row(cid: str) -> dict[str, Any]:
    c, r, w = CUST[cid], REN_BY_CUST[cid], WELL[cid]
    rk = RISK[r["renewal_id"]]
    top = NBA[cid][0]
    return {
        "customer_id": cid, "name": c["full_name"], "segment": c["segment"], "branch": c["branch_name"],
        "branch_id": c["branch_id"], "advisor": c["primary_advisor_name"], "renewal_id": r["renewal_id"],
        "days_to_maturity": int(r["days_to_maturity"]), "balance": float(r["mortgage_balance_cad"]),
        "payment_shock_pct": float(r["payment_shock_pct"]), "attrition_band": rk["attrition_risk_band"],
        "attrition_score": float(rk["attrition_risk_score"]), "lifetime_value": float(rk["five_year_lifetime_value_cad"]),
        "wellness_score": int(w["wellness_score"]), "wellness_band": w["wellness_band"],
        "tds_renewal_pct": float(CAP[cid]["tds_renewal_pct"]), "products_held": len(PRODUCTS[cid]),
        "top_action": {"type": top["action_type"], "product": top["product_name"], "eligibility": top["eligibility"]},
        "flags": _flags(cid), "email_consent": _consent(cid).get("Email", "Not captured"),
    }


def customers() -> dict[str, Any]:
    rows = [_row(cid) for cid in CUST]
    order = {"Vulnerable": 0, "Stretched": 1, "Stable": 2, "Thriving": 3}
    rows.sort(key=lambda x: (order[x["wellness_band"]], -x["attrition_score"]))
    return {"as_of": PB["as_of"], "customers": rows}


def overview() -> dict[str, Any]:
    rows = [_row(cid) for cid in CUST]
    bands = defaultdict(int)
    for x in rows:
        bands[x["wellness_band"]] += 1
    nba_types = defaultdict(int)
    blocked = defaultdict(int)
    for cid in CUST:
        for a in NBA[cid]:
            if a["eligibility"] == "Eligible":
                nba_types[a["action_type"]] += 1
            else:
                blocked["consent" if "consent" in a["eligibility"] else "vulnerable"] += 1
    return {
        "as_of": PB["as_of"],
        "kpis": {
            "customers": len(rows), "households": len(PB["households"]),
            "balance": round(sum(x["balance"] for x in rows)),
            "products_per_customer": round(sum(x["products_held"] for x in rows) / len(rows), 2),
            "vulnerable": bands["Vulnerable"], "stretched": bands["Stretched"],
            "deposit_flight": sum("Deposit flight" in x["flags"] for x in rows),
            "above_tds": sum(x["tds_renewal_pct"] > TDS_LIMIT for x in rows),
            "email_consent_pct": round(100 * sum(x["email_consent"] == "Granted" for x in rows) / len(rows)),
            "life_events": len(PB["life_events"]),
        },
        "wellness_bands": [{"band": b, "count": bands[b]} for b in ("Thriving", "Stable", "Stretched", "Vulnerable")],
        "actions": {"eligible": dict(nba_types), "blocked": dict(blocked)},
    }


def wellness() -> dict[str, Any]:
    rows = []
    for cid, w in WELL.items():
        if w["wellness_band"] not in ("Vulnerable", "Stretched"):
            continue
        sig = SIGNALS[cid]
        rows.append({**_row(cid), "drivers": [d for d in w["drivers"].split("; ") if d],
                     "support_pathway": w["support_pathway"], "sales_suppressed": w["sales_suppressed"] == "Y",
                     "strip": [{"month": s["month"], "overdraft_days": s["overdraft_days"], "nsf": s["nsf_count"],
                                "missed": s["missed_payments"], "util": s["card_utilization_pct"],
                                "payroll": s["payroll_inflow_cad"], "outflow": s["competitor_outflow_cad"]} for s in sig]})
    rows.sort(key=lambda x: x["wellness_score"])
    return {"as_of": PB["as_of"], "customers": rows,
            "policy": "RP-012 support-first; RP-014 hardship relief; RP-006 amortization fast path"}


def customer(cid: str) -> dict[str, Any] | None:
    if cid not in CUST:
        return None
    c, r, w, cap, hh = CUST[cid], REN_BY_CUST[cid], WELL[cid], CAP[cid], HH[cid]
    rk = RISK[r["renewal_id"]]
    consent = _consent(cid)
    return {
        **_row(cid),
        "household": hh, "province": c["province"], "tenure_years": int(c["tenure_years"]),
        "language": c["preferred_language"], "channel": c["origination_channel"], "nps": int(c["nps_score"]),
        "renewal": {**r, "attrition_risk_score": rk["attrition_risk_score"], "renewal_likelihood_pct": rk["renewal_likelihood_pct"],
                    "top_risk_drivers": rk["top_risk_drivers"], "competitor_rate_gap_bps": rk["competitor_rate_gap_bps"],
                    "recommended_offer_id": rk["recommended_offer_id"], "annual_revenue_exposure_cad": rk["annual_revenue_exposure_cad"]},
        "capacity": cap, "wellness": {**w, "drivers": [d for d in w["drivers"].split("; ") if d]},
        "products": PRODUCTS[cid], "signals": SIGNALS[cid], "life_events": EVENTS[cid],
        "consent": consent, "actions": NBA[cid],
        "prompts": _prompts(cid),
    }


def _prompts(cid: str) -> list[dict[str, str]]:
    c, r = CUST[cid], REN_BY_CUST[cid]
    rid = r["renewal_id"]
    return [
        {"label": "Brief me on this household", "prompt": f"Brief me on renewal {rid} and customer {cid} before my call: renewal terms, "
         "affordability at the renewal payment, financial-wellness signals, products held and consent. Verify in Fabric IQ and cite policy."},
        {"label": "What can I offer?", "prompt": f"For renewal {rid} ({cid}), which retention offers am I permitted to make under RP-001 to RP-003 "
         "and who approves, given the attrition score and competitor gap? Respect RP-012 if the client is vulnerable."},
        {"label": "Support or retention?", "prompt": f"Should {cid} (renewal {rid}) be handled as a retention case or a support-first case? "
         "Use financial_wellness, credit_capacity and RP-012 / RP-013 / RP-014, and list what is blocked and why."},
        {"label": "What did colleagues say?", "prompt": f"Search Work IQ for anything said about renewal {rid} or customer {c['full_name']} and summarise it."},
    ]


# --------------------------------------------------------------------------- #
# Outreach - deterministic draft + deterministic guardrails
# --------------------------------------------------------------------------- #

def _fmt_money(x: float) -> str:
    return f"${x:,.0f}"


def draft(cid: str, action_index: int = 0, channel: str = "Email", bank: str = "the Bank") -> dict[str, Any] | None:
    if cid not in CUST:
        return None
    c, r, w = CUST[cid], REN_BY_CUST[cid], WELL[cid]
    acts = NBA[cid]
    a = acts[max(0, min(action_index, len(acts) - 1))]
    first = c["full_name"].split()[0]
    maturity = date.fromisoformat(r["maturity_date"])
    term = r["product_type"].split(" Closed")[0]
    offer_by = (AS_OF + timedelta(days=30)).strftime("%B %d, %Y")
    if a["action_type"] == "Support":
        subject = "Planning ahead for your mortgage renewal"
        body = (f"Hi {first},\n\nYour mortgage renews on {maturity:%B %d, %Y}, and we know a new payment can take some planning. "
                "We'd like to set aside time to walk through your options together - including ways to make the payment more "
                "manageable, such as adjusting your amortization or a temporary payment arrangement, where you qualify.\n\n"
                "There's no cost and no obligation, and nothing changes unless you decide it should. You can reply to this "
                "message or call your advisor to book a time that suits you.\n\n"
                "We also offer free one-on-one budgeting support if that would help.\n\n"
                f"Kind regards,\n{c['primary_advisor_name']}\nYour mortgage advisor")
        purpose = "Service"
    elif a["action_type"] == "Retain":
        subject = "Your mortgage renewal options"
        body = (f"Hi {first},\n\nYour {term.lower()} mortgage renews on {maturity:%B %d, %Y}. We'd like to review your renewal options "
                "with you and make sure your next term fits your plans.\n\n"
                "Any rate we discuss will be for a stated term, subject to approval, and valid until "
                f"{offer_by}. We'll show you the payment, amortization and frequency it assumes before you decide.\n\n"
                "Reply to this message or call me to book a time.\n\n"
                f"Kind regards,\n{c['primary_advisor_name']}\nYour mortgage advisor")
        purpose = "Service"
    else:
        subject = f"An idea for you: {a['product_name']}"
        body = (f"Hi {first},\n\nAs you plan for your renewal on {maturity:%B %d, %Y}, you may want to consider a "
                f"{a['product_name']}. {_why(a)}\n\n"
                "If you're interested, I can walk you through how it works, the fees and whether it suits your situation. "
                "Any product is subject to approval.\n\n"
                f"Kind regards,\n{c['primary_advisor_name']}\nYour mortgage advisor\n\n"
                f"You are receiving this because you agreed to hear from {bank}. To stop receiving these messages, reply STOP or use the unsubscribe link.")
        purpose = "Marketing"
    d = {"customer_id": cid, "action": a, "channel": channel, "purpose": purpose, "subject": subject, "body": body}
    return {**d, "checks": checks(cid, d)}


def _why(a: dict) -> str:
    return {"PRD-TFSA": "It's a flexible way to grow savings tax-free.",
            "PRD-RSP": "It can help you save for retirement with tax-deductible contributions.",
            "PRD-HEL": "It can give you flexible access to funds for planned expenses, secured against your home.",
            "PRD-INS": "It can help protect your family's home if the unexpected happens.",
            "PRD-CRD": "It brings everyday spending and rewards together with your other accounts."}.get(a["product_id"], "")


def checks(cid: str, d: dict[str, Any]) -> list[dict[str, str]]:
    """Deterministic pre-send guardrails. status: pass | warn | block."""
    w, cap = WELL[cid], CAP[cid]
    text = f"{d.get('subject', '')}\n{d.get('body', '')}"
    channel = d.get("channel", "Email")
    purpose = d.get("purpose", "Service")
    promo = purpose == "Marketing" or bool(re.search(r"\b(consider a|offer|special|promotion|bonus|rewards)\b", text, re.I)) and purpose != "Service"
    out = []
    consent = _consent(cid).get(channel, "Not captured")
    if purpose == "Marketing":
        out.append({"rule": "RP-011", "name": "Marketing consent", "status": "pass" if consent == "Granted" else "block",
                    "detail": f"{channel} marketing consent is {consent}."})
        unsub = bool(re.search(r"unsubscribe|reply stop", text, re.I))
        out.append({"rule": "RP-011", "name": "Unsubscribe mechanism", "status": "pass" if unsub else "block",
                    "detail": "Unsubscribe present." if unsub else "Promotional message has no unsubscribe mechanism."})
    else:
        sales = re.search(r"\b(TFSA|RRSP|retirement savings|credit card|HELOC|line of credit|insurance|investment)\b", text, re.I)
        out.append({"rule": "RP-011", "name": "Service message stays service", "status": "block" if sales else "pass",
                    "detail": f"Service message contains a product promotion ('{sales.group(0)}')." if sales else "No promotional call to action."})
    vul = w["vulnerability_flag"] == "Y"
    out.append({"rule": "RP-012", "name": "Vulnerable client", "status": ("block" if vul and purpose == "Marketing" else "warn" if vul else "pass"),
                "detail": ("Sales are suppressed for this client - support-first only." if vul and purpose == "Marketing" else
                           "Support-first client: Branch Manager approval required before sending." if vul else "No vulnerability flag.")})
    tds = float(cap["tds_renewal_pct"])
    credit = d.get("action", {}).get("product_id") in ("PRD-HEL", "PRD-CRD")
    out.append({"rule": "RP-013", "name": "Affordability", "status": "block" if credit and tds > TDS_LIMIT else "pass",
                "detail": (f"New credit cannot be recommended: renewal TDS {tds}% exceeds the {TDS_LIMIT:.0f}% guideline." if credit and tds > TDS_LIMIT
                           else f"Renewal TDS {tds}%.")})
    hits = []
    for pat, rule, why in BANNED:
        m = re.search(pat, text, re.I)
        if m:
            hits.append({"rule": rule, "name": why, "status": "block", "detail": f"Found: \"{m.group(0)}\""})
    out.extend(hits or [{"rule": "RP-016", "name": "Prohibited words and internal data", "status": "pass", "detail": "None found."}])
    if re.search(r"\b\d+(\.\d+)?\s?%", text) and not re.search(r"subject to approval", text, re.I):
        out.append({"rule": "RP-016", "name": "Rate disclosure", "status": "block", "detail": "A rate is quoted without 'subject to approval'."})
    del promo
    return out


def verdict(cks: list[dict[str, str]]) -> str:
    return "blocked" if any(c["status"] == "block" for c in cks) else "needs_approval" if any(c["status"] == "warn" for c in cks) else "pass"


def review_prompt(cid: str, d: dict[str, Any], cks: list[dict[str, str]]) -> str:
    w = WELL[cid]
    det = "\n".join(f"- {c['rule']} {c['name']}: {c['status'].upper()} - {c['detail']}" for c in cks)
    return (
        f"Review this client communication before it can be saved as a draft. Client context (internal, do NOT repeat to the client): "
        f"wellness band {w['wellness_band']}, sales suppressed {w['sales_suppressed']}, channel {d.get('channel')}, purpose {d.get('purpose')}, "
        f"next-best action {d.get('action', {}).get('action_type')} / {d.get('action', {}).get('product_name')}.\n"
        f"Deterministic checks already run (treat as authoritative):\n{det}\n\n"
        f"DRAFT SUBJECT: {d.get('subject', '')}\nDRAFT BODY:\n{d.get('body', '')}\n\n"
        "Return: (1) a verdict line 'VERDICT: PASS', 'VERDICT: NEEDS FIXES' or 'VERDICT: BLOCKED' - never weaker than the deterministic "
        "checks; (2) a table of findings with the exact offending text, the policy id (RP-011 to RP-018, RP-007) and why; (3) a corrected "
        "draft if fixes are possible. Judge tone (empathetic, non-coercive), implied promises, fairness and privacy. Refer to the firm as "
        "'the Bank'. You review; you never send or approve."
    )


# --------------------------------------------------------------------------- #
# Wrap-up memo (.docx)
# --------------------------------------------------------------------------- #

def memo_docx(cid: str, d: dict[str, Any], notes: str, upn: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    c = customer(cid)
    assert c
    doc = Document()
    doc.styles["Normal"].font.name = "Segoe UI"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(f"Client conversation note - {c['renewal_id']}", level=1)
    doc.add_paragraph("SYNTHETIC DEMO DATA. Internal record - not for the client.").italic = True
    t = doc.add_table(rows=0, cols=2)
    for k, v in [("Client", c["name"]), ("Household", c["household"]["household_name"]), ("Advisor", c["advisor"]),
                 ("Renewal", f"{c['renewal']['product_type']} - matures {c['renewal']['maturity_date']}"),
                 ("Wellness band", c["wellness_band"]), ("Next-best action", f"{d['action']['action_type']} - {d['action']['product_name']}"),
                 ("Prepared by", upn), ("As of", PB["as_of"])]:
        row = t.add_row().cells
        row[0].text, row[1].text = k, str(v)
    doc.add_heading("Facts from the system of record (Fabric IQ)", level=2)
    for line in [f"Renewal payment {_fmt_money(float(c['capacity']['renewal_payment_monthly_cad']))}/month "
                 f"({c['payment_shock_pct']:.0f}% above today's)", f"TDS at renewal {c['tds_renewal_pct']}% (guideline {TDS_LIMIT:.0f}%)",
                 f"Products held: {', '.join(p['product_name'] for p in c['products'])}"] + c["wellness"]["drivers"]:
        doc.add_paragraph(line, style="List Bullet")
    doc.add_heading("What the client said", level=2)
    doc.add_paragraph(notes or "(not recorded)")
    doc.add_heading("Outbound draft (not sent)", level=2)
    doc.add_paragraph(f"Subject: {d['subject']}")
    doc.add_paragraph(d["body"])
    doc.add_paragraph("AI-assisted content above is labelled as a draft under RP-017; nothing was sent automatically.").italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
